"""Recalculate formula workbooks without changing the uploaded original."""

from __future__ import annotations

import os
import shutil
import signal
import struct
import subprocess
from contextlib import closing
from pathlib import Path
from tempfile import TemporaryDirectory

from octop.i18n import tr
from octop.infra.errors import ErrorCode, OctopError

SPREADSHEET_SUFFIXES = frozenset({".xlsx", ".xlsm", ".xls"})
RECALC_TIMEOUT_SECONDS = 120
# A private profile prevents attachment to a user's existing LibreOffice process.
_PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">
 <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
</item>
<item oor:path="/org.openoffice.Office.Calc/Content/Update">
 <prop oor:name="Link" oor:op="fuse"><value>1</value></prop>
</item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting">
 <prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>
 <prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop>
 <prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop>
</item>
</oor:items>"""


def _xlsx_formulas(path: Path) -> tuple[list[str], set[tuple[str, str]]]:
    from openpyxl import load_workbook

    with closing(load_workbook(path, read_only=True, data_only=False)) as book:
        return book.sheetnames, {
            (sheet.title, cell.coordinate)
            for sheet in book
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }


def _xls_formulas(path: Path) -> tuple[list[str], int]:
    import xlrd
    from xlrd.compdoc import CompDoc

    data = path.read_bytes()
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        compound = CompDoc(data)
        data = compound.get_named_stream("Workbook") or compound.get_named_stream("Book")
    with xlrd.open_workbook(path, on_demand=True) as book:
        names = book.sheet_names()
    count = 0
    offset = 0
    # BIFF FORMULA records retain formula presence even when xlrd returns only cached values.
    while offset + 4 <= len(data):
        opcode, size = struct.unpack_from("<HH", data, offset)
        if offset + 4 + size > len(data):
            raise ValueError("truncated BIFF record")
        count += opcode in {0x0006, 0x0206, 0x0406}
        offset += 4 + size
    return names, count


def spreadsheet_preview_path(original: Path, derived: Path, status: str) -> Path:
    if original.suffix.lower() not in SPREADSHEET_SUFFIXES:
        return original
    if status != "ready":
        raise OctopError.localized(ErrorCode.KNOWLEDGE_SPREADSHEET_NOT_READY)
    if derived.is_file():
        return derived
    _, formulas = (
        _xls_formulas(original) if original.suffix.lower() == ".xls" else _xlsx_formulas(original)
    )
    if formulas:
        raise OctopError.localized(ErrorCode.KNOWLEDGE_SPREADSHEET_NOT_READY)
    return original


def prepare_spreadsheet(original: Path, derived: Path, *, locale: str = "en") -> Path:
    """Return the original for static sheets, otherwise publish a validated XLSX copy."""
    if original.suffix.lower() not in SPREADSHEET_SUFFIXES:
        return original
    derived.unlink(missing_ok=True)
    names, formulas = (
        _xls_formulas(original) if original.suffix.lower() == ".xls" else _xlsx_formulas(original)
    )
    if not formulas:
        return original
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if binary is None:
        raise RuntimeError(tr("knowledge.spreadsheet_unavailable", locale))
    try:
        with TemporaryDirectory(prefix="excel-recalc-", dir=original.parent) as temp:
            work = Path(temp)
            profile = work / "profile"
            (profile / "user").mkdir(parents=True)
            (profile / "user" / "registrymodifications.xcu").write_text(_PROFILE, encoding="utf-8")
            output = work / "output"
            output.mkdir()
            with subprocess.Popen(
                [
                    binary,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML",
                    "--outdir",
                    str(output),
                    str(original.resolve()),
                ],
                stdin=subprocess.DEVNULL,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                start_new_session=os.name == "posix",
            ) as process:
                try:
                    code = process.wait(timeout=RECALC_TIMEOUT_SECONDS)
                except subprocess.TimeoutExpired:
                    if os.name == "posix":
                        os.killpg(process.pid, signal.SIGKILL)
                    else:
                        try:
                            subprocess.run(
                                ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                                stdout=subprocess.DEVNULL,
                                stderr=subprocess.DEVNULL,
                                timeout=5,
                                check=False,
                            )
                        finally:
                            process.kill()
                    process.wait()
                    raise
            result = output / (original.stem + ".xlsx")
            if code != 0 or not result.is_file():
                raise ValueError("LibreOffice produced no workbook")
            result_names, result_formulas = _xlsx_formulas(result)
            if result_names != names or (
                len(result_formulas) < formulas
                if isinstance(formulas, int)
                else not formulas.issubset(result_formulas)
            ):
                raise ValueError("LibreOffice changed sheets or removed formulas")
            from openpyxl import load_workbook

            with closing(load_workbook(result, read_only=True, data_only=True)) as book:
                for sheet in book:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.data_type == "e" or (
                                (sheet.title, getattr(cell, "coordinate", "")) in result_formulas
                                and cell.value is None
                                and cell.data_type not in {"str", "s", "inlineStr"}
                            ):
                                raise ValueError(
                                    f"invalid formula result: {sheet.title}!{cell.coordinate}"
                                )
            if not original.is_file():
                raise FileNotFoundError(original)
            result.replace(derived)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(tr("knowledge.spreadsheet_timeout", locale)) from exc
    except Exception as exc:
        raise RuntimeError(tr("knowledge.spreadsheet_failed", locale)) from exc
    return derived
