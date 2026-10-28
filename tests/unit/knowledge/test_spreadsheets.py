"""Formula detection, recalculation failures, and one-record knowledge previews."""

from __future__ import annotations

import io
import os
import subprocess
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from octop.infra.db.migrate import run_migrations
from octop.infra.db.pool import SqlitePool
from octop.infra.db.repos.knowledge import KnowledgeRepo
from octop.infra.db.repos.settings import SettingsRepo
from octop.infra.db.repos.users import UserRepo
from octop.infra.errors import ErrorCode, OctopError
from octop.infra.knowledge import jobs, spreadsheets
from octop.infra.knowledge import service as service_module
from octop.infra.knowledge.files import document_path, recalculated_document_path
from octop.infra.knowledge.index import KnowledgeIndex
from octop.infra.knowledge.service import KnowledgeService


def workbook_bytes(*, formula: bool = True, cache: str = "") -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Amounts"
    sheet.append(["Item", "Amount"])
    sheet.append(["Revenue", "=SUM(C2:D2)" if formula else 125.5, 100, 25.5])
    sheet["A4"] = "=literal text"
    sheet["A4"].data_type = "s"
    raw = io.BytesIO()
    book.save(raw)
    book.close()
    if not cache:
        return raw.getvalue()
    cached = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(cached, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b"<f>SUM(C2:D2)</f><v></v>", f"<f>SUM(C2:D2)</f>{cache}".encode()
                )
                if cache == "<v/>" or cache == "<v>#VALUE!</v>":
                    kind = b"str" if cache == "<v/>" else b"e"
                    data = data.replace(b'<c r="B2">', b'<c r="B2" t="' + kind + b'">')
            target.writestr(info, data)
    return cached.getvalue()


def mock_converter(monkeypatch: pytest.MonkeyPatch, result: bytes | None) -> MagicMock:
    process = MagicMock()
    process.__enter__.return_value = process
    process.wait.return_value = 0
    process.pid = 12345

    def start(args: list[str], **kwargs: object) -> MagicMock:
        out = Path(args[args.index("--outdir") + 1])
        profile = out.parent / "profile" / "user" / "registrymodifications.xcu"
        assert "OOXMLRecalcMode" in profile.read_text()
        assert "DisableMacrosExecution" in profile.read_text()
        assert "BlockUntrustedRefererLinks" in profile.read_text()
        if result is not None:
            (out / (Path(args[-1]).stem + ".xlsx")).write_bytes(result)
        return process

    monkeypatch.setattr(spreadsheets.shutil, "which", lambda _: "libreoffice-test")
    monkeypatch.setattr(spreadsheets.subprocess, "Popen", start)
    return process


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm", ".xls", ".csv"])
def test_static_sheets_never_start_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, suffix: str
) -> None:
    path = tmp_path / ("static" + suffix)
    if suffix == ".xls":
        import xlwt

        book = xlwt.Workbook()
        book.add_sheet("Amounts").write(0, 0, 125.5)
        book.save(str(path))
    elif suffix == ".csv":
        path.write_text("Amount\n125.5")
    else:
        path.write_bytes(workbook_bytes(formula=False))
    original = path.read_bytes()
    run = MagicMock(side_effect=AssertionError("must not start LibreOffice"))
    monkeypatch.setattr(spreadsheets.subprocess, "Popen", run)
    monkeypatch.setattr(spreadsheets.shutil, "which", lambda _: None)
    assert spreadsheets.prepare_spreadsheet(path, tmp_path / "copy.xlsx") == path
    assert path.read_bytes() == original
    run.assert_not_called()


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm", ".xls"])
@pytest.mark.parametrize("cache", ["<v>125.5</v>", "<v>0</v>", "<v/>"])
def test_formulas_publish_validated_copy(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, suffix: str, cache: str
) -> None:
    path, derived = tmp_path / ("amounts" + suffix), tmp_path / "copy.xlsx"
    if suffix == ".xls":
        import xlwt

        book = xlwt.Workbook()
        book.add_sheet("Amounts").write(1, 1, xlwt.Formula("SUM(C2:D2)"))
        book.save(str(path))
    else:
        path.write_bytes(workbook_bytes())
    original = path.read_bytes()
    expected = workbook_bytes(cache=cache)
    process = mock_converter(monkeypatch, expected)
    assert spreadsheets.prepare_spreadsheet(path, derived) == derived
    assert derived.read_bytes() == expected
    assert path.read_bytes() == original
    process.wait.assert_called_once_with(timeout=120)
    assert not list(tmp_path.glob("excel-recalc-*"))


@pytest.mark.parametrize(
    "result",
    [
        None,
        b"bad zip",
        workbook_bytes(),
        workbook_bytes(cache="<v>#VALUE!</v>"),
        workbook_bytes(formula=False),
    ],
    ids=["no-output", "invalid-zip", "missing-cache", "formula-error", "missing-formula"],
)
def test_failed_conversion_does_not_publish_or_keep_stale_copy(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, result: bytes | None
) -> None:
    path, derived = tmp_path / "amounts.xlsx", tmp_path / "copy.xlsx"
    path.write_bytes(workbook_bytes())
    derived.write_bytes(b"old copy")
    original = path.read_bytes()
    mock_converter(monkeypatch, result)
    with pytest.raises(RuntimeError, match="recalculation failed"):
        spreadsheets.prepare_spreadsheet(path, derived)
    assert not derived.exists()
    assert path.read_bytes() == original
    assert not list(tmp_path.glob("excel-recalc-*"))


def test_missing_libreoffice_is_actionable(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "amounts.xlsx"
    path.write_bytes(workbook_bytes())
    monkeypatch.setattr(spreadsheets.shutil, "which", lambda _: None)
    with pytest.raises(RuntimeError, match="Install LibreOffice Calc"):
        spreadsheets.prepare_spreadsheet(path, tmp_path / "copy.xlsx")


def test_timeout_kills_converter_and_cleans_staging(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, derived = tmp_path / "amounts.xlsx", tmp_path / "copy.xlsx"
    path.write_bytes(workbook_bytes())
    process = mock_converter(monkeypatch, workbook_bytes(cache="<v>125.5</v>"))
    process.wait.side_effect = [subprocess.TimeoutExpired("soffice", 120), 0]
    killpg = MagicMock()
    taskkill = MagicMock()
    monkeypatch.setattr(spreadsheets.subprocess, "run", taskkill)
    if os.name == "posix":
        monkeypatch.setattr(spreadsheets.os, "killpg", killpg)
    with pytest.raises(RuntimeError, match="timed out"):
        spreadsheets.prepare_spreadsheet(path, derived)
    if os.name == "posix":
        killpg.assert_called_once()
    else:
        process.kill.assert_called_once()
        taskkill.assert_called_once()
        assert "/T" in taskkill.call_args.args[0]
    assert not derived.exists()
    assert not list(tmp_path.glob("excel-recalc-*"))


@pytest.fixture
def knowledge(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "db.sqlite")
    run_migrations(pool)
    services = SimpleNamespace(
        knowledge_repo=KnowledgeRepo(pool),
        settings_repo=SettingsRepo(pool),
        user_repo=UserRepo(pool),
    )
    owner = services.user_repo.create(username="owner", password_hash="h", role="user")
    base = services.knowledge_repo.create_base(
        owner_user_id=owner, name="Reports", embedding_model="test"
    )
    monkeypatch.setattr(jobs, "assert_knowledge_usable", lambda *_: None)
    monkeypatch.setattr(service_module, "assert_knowledge_usable", lambda *_: None)
    monkeypatch.setattr(
        jobs, "embed_knowledge_texts", lambda _services, texts: [[1.0, 0.0] for _ in texts]
    )
    service = KnowledgeService(services)
    doc = service.upload_document(
        base.id,
        actor_user_id=owner,
        filename="report.xlsx",
        content_type="",
        content=workbook_bytes(),
    )
    return services, service, owner, base.id, doc.id


def test_index_and_previews_share_copy_download_keeps_original_one_row(
    knowledge, monkeypatch: pytest.MonkeyPatch
) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    original = document_path(kb_id, doc_id, "report.xlsx")
    derived = recalculated_document_path(kb_id, doc_id)
    before = original.read_bytes()
    with pytest.raises(OctopError) as exc:
        service.resolve_document_file(kb_id, doc_id, actor_user_id=owner, for_preview=True)
    assert exc.value.code == ErrorCode.KNOWLEDGE_SPREADSHEET_NOT_READY
    mock_converter(monkeypatch, workbook_bytes(cache="<v>125.5</v>"))
    jobs.process_document(services, kb_id, doc_id)
    assert services.knowledge_repo.get_document(doc_id).status == "ready"
    assert len(service.list_documents(kb_id, actor_user_id=owner)) == 1
    assert service.resolve_document_file(kb_id, doc_id, actor_user_id=owner)[0] == original
    assert (
        service.resolve_document_file(kb_id, doc_id, actor_user_id=owner, for_preview=True)[0]
        == derived
    )
    assert "125.5" in service.preview_document(kb_id, doc_id, actor_user_id=owner)["text"]
    assert "125.5" in KnowledgeIndex(kb_id).search([1.0, 0.0], 1)[0].text
    assert original.read_bytes() == before
    service.reindex_document(kb_id, doc_id, actor_user_id=owner)
    jobs.process_document(services, kb_id, doc_id)
    assert len(service.list_documents(kb_id, actor_user_id=owner)) == 1
    service.delete_document(kb_id, doc_id, actor_user_id=owner)
    assert not original.exists() and not derived.exists()
    assert not KnowledgeIndex(kb_id).search([1.0, 0.0], 1)


def test_failed_recalculation_marks_failed_without_indexing(
    knowledge, monkeypatch: pytest.MonkeyPatch
) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    mock_converter(monkeypatch, workbook_bytes())
    embed = MagicMock(side_effect=AssertionError("bad amounts must not be embedded"))
    monkeypatch.setattr(jobs, "embed_knowledge_texts", embed)
    with pytest.raises(RuntimeError):
        jobs.process_document(services, kb_id, doc_id)
    row = services.knowledge_repo.get_document(doc_id)
    assert row.status == "failed" and row.chunk_count == 0
    assert "重算失败" in row.error_message
    assert not recalculated_document_path(kb_id, doc_id).exists()
    assert service.resolve_document_file(kb_id, doc_id, actor_user_id=owner)[0].exists()
    with pytest.raises(OctopError):
        service.preview_document(kb_id, doc_id, actor_user_id=owner)
    embed.assert_not_called()


def test_ready_formula_document_without_copy_requires_reindex(knowledge) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    services.knowledge_repo.update_document(doc_id, status="ready")
    with pytest.raises(OctopError):
        service.resolve_document_file(kb_id, doc_id, actor_user_id=owner, for_preview=True)
    assert service.resolve_document_file(kb_id, doc_id, actor_user_id=owner)[0].exists()


def test_ready_static_sheet_preview_uses_original(knowledge) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    original = document_path(kb_id, doc_id, "report.xlsx")
    original.write_bytes(workbook_bytes(formula=False))
    stale = recalculated_document_path(kb_id, doc_id)
    stale.write_bytes(b"old copy")
    jobs.process_document(services, kb_id, doc_id)
    assert not stale.exists()
    assert (
        service.resolve_document_file(kb_id, doc_id, actor_user_id=owner, for_preview=True)[0]
        == original
    )


def test_existing_formula_cache_still_recalculates(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original, derived = tmp_path / "cached.xlsx", tmp_path / "copy.xlsx"
    original.write_bytes(workbook_bytes(cache="<v>999</v>"))
    process = mock_converter(monkeypatch, workbook_bytes(cache="<v>125.5</v>"))
    assert spreadsheets.prepare_spreadsheet(original, derived) == derived
    process.wait.assert_called_once()
    assert derived.read_bytes() != original.read_bytes()


def test_delete_during_indexing_cleans_copy(knowledge, monkeypatch: pytest.MonkeyPatch) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    mock_converter(monkeypatch, workbook_bytes(cache="<v>125.5</v>"))

    def embed(_services, texts):
        service.delete_document(kb_id, doc_id, actor_user_id=owner)
        return [[1.0, 0.0] for _ in texts]

    monkeypatch.setattr(jobs, "embed_knowledge_texts", embed)
    jobs.process_document(services, kb_id, doc_id)
    assert not recalculated_document_path(kb_id, doc_id).exists()
    assert not KnowledgeIndex(kb_id).search([1.0, 0.0], 1)


def test_preview_copy_preserves_read_access(knowledge, monkeypatch: pytest.MonkeyPatch) -> None:
    services, service, owner, kb_id, doc_id = knowledge
    reader = services.user_repo.create(username="reader", password_hash="h", role="user")
    service.update_base(kb_id, actor_user_id=owner, shared=True)
    mock_converter(monkeypatch, workbook_bytes(cache="<v>125.5</v>"))
    jobs.process_document(services, kb_id, doc_id)
    assert "125.5" in service.preview_document(kb_id, doc_id, actor_user_id=reader)["text"]
    assert service.resolve_document_file(kb_id, doc_id, actor_user_id=reader, for_preview=True)[
        0
    ] == recalculated_document_path(kb_id, doc_id)
    service.update_base(kb_id, actor_user_id=owner, shared=False)
    with pytest.raises(PermissionError):
        service.resolve_document_file(kb_id, doc_id, actor_user_id=reader, for_preview=True)
