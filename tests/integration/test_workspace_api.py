"""tests/integration/test_workspace_api.py — workspace endpoints.

Requires a running harness agent (``env_with_agent``); workspace I/O goes
through ``agent.workspace`` backed by ``local_shell`` on the agent dir.
"""

from __future__ import annotations

import zipfile
from io import BytesIO
from typing import Any

import pytest
from docx import Document

from octop.api.common.agent_workspace import resolve_agent_workspace_dir

# Workspace UI semantics: leading '/' is relative to agent workspace.
FROM_WORKSPACE = {"from_workspace": "true"}


def _sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Report Title", level=1)
    doc.add_paragraph("Intro paragraph")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
async def env(env_with_agent):
    yield env_with_agent


# --- listing ---------------------------------------------------------------


async def test_tree_returns_empty_for_fresh_workspace(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(f"/api/agents/{aid}/workspace/tree?from_workspace=true", headers=auth)
    assert r.status_code == 200, r.text
    rows = r.json()
    assert isinstance(rows, list)
    # Fresh workspace may contain a SOUL.md (written at agent boot) or
    # be empty — we don't pin the exact contents, just the shape.
    for row in rows:
        assert "path" in row


async def test_tree_lists_root_files(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("notes.md", b"hello")

    r = await c.get(f"/api/agents/{aid}/workspace/tree?path=/&from_workspace=true", headers=auth)
    assert r.status_code == 200, r.text
    rows = r.json()
    paths = {row["path"] for row in rows}
    assert any("notes.md" in p for p in paths)

    r = await c.get(
        f"/api/agents/{aid}/workspace/file?path=%2Fnotes.md&from_workspace=true",
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "hello"


async def test_tree_lists_subdirectory(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/sub/nested.txt"},
        headers=auth,
        json={"content": "nested content"},
    )

    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/sub"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    rows = r.json()
    assert any("nested.txt" in row["path"] for row in rows)

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/sub/nested.txt"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "nested content"


async def test_tree_for_unknown_agent_404(env: Any) -> None:
    c, _srv, auth, _aid = env
    r = await c.get("/api/agents/no-such-agent/workspace/tree?from_workspace=true", headers=auth)
    assert r.status_code == 404


# --- write + read round-trip ------------------------------------------------


async def test_write_then_read_roundtrip(env: Any) -> None:
    c, _srv, auth, aid = env
    payload = "hello from workspace test\n"
    r = await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
        json={"content": payload},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/notes.md"
    assert r.json()["size"] == len(payload.encode("utf-8"))

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
    )
    assert r.status_code == 200
    body = r.json()
    assert body["path"] == "/notes.md"
    assert body["content"] == payload


async def test_read_missing_file_404(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/no-such-file.txt"},
        headers=auth,
    )
    assert r.status_code == 404


# --- upload + download ------------------------------------------------------


async def test_upload_then_download_binary(env: Any) -> None:
    c, _srv, auth, aid = env
    blob = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16  # fake PNG header
    files = {"file": ("logo.png", blob, "image/png")}
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE},
        headers=auth,
        files=files,
    )
    assert r.status_code == 200, r.text
    assert r.json()["size"] == len(blob)

    r = await c.get(
        f"/api/agents/{aid}/workspace/download",
        params={**FROM_WORKSPACE, "path": "/logo.png"},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.content.startswith(b"\x89PNG\r\n\x1a\n")
    cd = r.headers.get("content-disposition", "")
    assert "logo.png" in cd


async def test_download_non_ascii_filename(env: Any) -> None:
    c, _srv, auth, aid = env
    fname = "1783510288_地球介绍.pptx"
    path = f"/outbound/{fname}"
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE, "path": path},
        headers=auth,
        files={"file": (fname, b"PK\x03\x04fake", "application/vnd.ms-powerpoint")},
    )
    assert r.status_code == 200, r.text

    r = await c.get(
        f"/api/agents/{aid}/workspace/download",
        params={**FROM_WORKSPACE, "path": path},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"PK\x03\x04")
    cd = r.headers.get("content-disposition", "")
    assert 'filename="download.pptx"' in cd
    assert "filename*" in cd
    assert "%E5%9C%B0%E7%90%83" in cd


async def test_upload_with_explicit_path_query(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE, "path": "/sub/dir/named.txt"},
        headers=auth,
        files={"file": ("ignored.txt", b"x", "text/plain")},
    )
    assert r.status_code == 200
    assert r.json()["path"] == "/sub/dir/named.txt"


# --- glob + grep ------------------------------------------------------------


async def test_glob_after_seeding(env: Any) -> None:
    c, _srv, auth, aid = env
    for fname in ("a.md", "b.md", "c.txt"):
        await c.put(
            f"/api/agents/{aid}/workspace/file",
            params={**FROM_WORKSPACE, "path": f"/{fname}"},
            headers=auth,
            json={"content": "x"},
        )
    r = await c.get(
        f"/api/agents/{aid}/workspace/glob",
        params={**FROM_WORKSPACE, "pattern": "*.md", "path": "/"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    paths = {row["path"] for row in r.json()}
    # Glob may return absolute or relative paths depending on backend.
    matched = {p.rsplit("/", 1)[-1] for p in paths}
    assert "a.md" in matched
    assert "b.md" in matched


async def test_grep_after_seeding(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/needle.txt"},
        headers=auth,
        json={"content": "alpha\nNEEDLE here\ngamma\n"},
    )
    r = await c.get(
        f"/api/agents/{aid}/workspace/grep",
        params={**FROM_WORKSPACE, "pattern": "NEEDLE", "path": "/"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    rows = r.json()
    assert any("NEEDLE" in str(row) for row in rows)


# --- cross-user isolation ---------------------------------------------------


async def test_non_owner_cannot_access_workspace(env: Any) -> None:
    """Non-owners cannot read another user's agent workspace."""
    c, _srv, admin_auth, _aid = env
    await c.post(
        "/api/users",
        headers=admin_auth,
        json={"username": "bob", "password": "TestPass12", "role": "user"},
    )
    bob_tok = (
        await c.post(
            "/api/auth/login",
            json={"username": "bob", "password": "TestPass12"},
        )
    ).json()["access_token"]
    bob_auth = {"Authorization": f"Bearer {bob_tok}"}

    admin_agent_id = (await c.get("/api/agents", headers=admin_auth)).json()[0]["agent_id"]

    r = await c.get(
        f"/api/agents/{admin_agent_id}/workspace/tree",
        headers=bob_auth,
    )
    assert r.status_code == 403


# --- delete + move ----------------------------------------------------------


async def test_delete_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
        json={"content": "bye"},
    )
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
    )
    assert r.status_code == 204, r.text

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_move_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
        json={"content": "payload"},
    )
    r = await c.post(
        f"/api/agents/{aid}/workspace/move",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
        json={"destination": "/moved/src.txt"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/moved/src.txt"

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/moved/src.txt"},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.json()["content"] == "payload"

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_rename_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/old-name.md"},
        headers=auth,
        json={"content": "x"},
    )
    r = await c.post(
        f"/api/agents/{aid}/workspace/move",
        params={**FROM_WORKSPACE, "path": "/old-name.md"},
        headers=auth,
        json={"destination": "/new-name.md"},
    )
    assert r.status_code == 200, r.text
    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
    )
    paths = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "new-name.md" in paths
    assert "old-name.md" not in paths


async def test_mkdir_creates_directory(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.post(
        f"/api/agents/{aid}/workspace/mkdir",
        params={**FROM_WORKSPACE, "path": "/projects/demo"},
        headers=auth,
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["path"] == "/projects/demo"
    assert body["is_dir"] is True

    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/projects"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    names = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "demo" in names


async def test_delete_directory(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/box/a.txt"},
        headers=auth,
        json={"content": "a"},
    )
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/box"},
        headers=auth,
    )
    assert r.status_code == 204, r.text
    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
    )
    names = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "box" not in names


async def test_delete_builtin_skills_forbidden(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/_builtin_skills/foo/SKILL.md"},
        headers=auth,
    )
    assert r.status_code == 403


# --- editable document (Markdown round-trip) --------------------------------


async def test_doc_read_and_write_roundtrip(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("report.docx", _sample_docx_bytes())

    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["path"] == "/report.docx"
    assert "# Report Title" in body["content"]
    assert "Intro paragraph" in body["content"]

    r = await c.put(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
        json={"content": "# Updated Title\n\nNew **bold** body\n"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/report.docx"
    assert r.json()["size"] > 0

    # Stored bytes parse back into a valid docx.
    blob = await agent.workspace.adownload_bytes("report.docx")
    assert blob is not None
    parsed = Document(BytesIO(blob))
    texts = [p.text for p in parsed.paragraphs if p.text]
    assert "Updated Title" in texts
    assert "New bold body" in texts

    # Round-trip back through the API keeps the Markdown structure.
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    content = r.json()["content"]
    assert "# Updated Title" in content
    assert "**bold**" in content


async def test_doc_unsupported_extension_400(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
    )
    assert r.status_code == 400


async def test_create_empty_docx_via_text_endpoint_is_previewable(env: Any) -> None:
    """Workspace "new file" creates .docx through the text endpoint; it must be
    stored as a valid document package so preview/edit work immediately."""
    c, srv, auth, aid = env
    r = await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/fresh.docx"},
        headers=auth,
        json={"content": ""},
    )
    assert r.status_code == 200, r.text
    assert r.json()["size"] > 0  # not a 0-byte file

    agent = srv.app_runtime.agent_registry.get_agent(aid)
    blob = await agent.workspace.adownload_bytes("fresh.docx")
    assert blob is not None
    parsed = Document(BytesIO(blob))
    assert len(parsed.paragraphs) == 0  # valid empty document

    # It can be opened for editing as an empty Markdown document.
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/fresh.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == ""


async def test_doc_missing_file_404(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/no-such.docx"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_doc_write_root_forbidden(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.put(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
        json={"content": "# hi\n"},
    )
    assert r.status_code == 403


async def test_doc_write_invalid_content_400(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("broken.docx", b"not a real docx zip")

    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/broken.docx"},
        headers=auth,
    )
    assert r.status_code == 400


# --- B-19-A: team memory writes are denied (409) ----------------------------

TEAM_LEARNINGS = ".octop/team/LEARNINGS.md"
TEAM_MEMORY_DENIED_EN = "You are not the owner of this artifact."
TEAM_MEMORY_DENIED_ZH = "你不是该工件的归属写者。"
# Direct spelling plus the three bypass spellings that defeat a raw-prefix check.
TEAM_MEMORY_SPELLINGS = [
    "/.octop/team/LEARNINGS.md",
    "/./.octop/team/LEARNINGS.md",  # ./ bypass
    "/x/../.octop/team/LEARNINGS.md",  # x/../ bypass
    "//.octop/team/LEARNINGS.md",  # // bypass
]


def _zip_with(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, blob in entries.items():
            archive.writestr(name, blob)
    return buffer.getvalue()


async def _team_memory_write(
    c: Any, auth: dict[str, str], aid: str, *, method: str, path: str, locale: str
) -> Any:
    headers = {**auth, "Accept-Language": locale}
    params = {**FROM_WORKSPACE, "path": path}
    if method == "put":
        return await c.put(
            f"/api/agents/{aid}/workspace/file",
            params=params,
            headers=headers,
            json={"content": "forged\n"},
        )
    return await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params=params,
        headers=headers,
        files={"file": ("LEARNINGS.md", b"forged\n", "text/markdown")},
    )


async def _assert_not_on_disk(c: Any, auth: dict[str, str], aid: str) -> None:
    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": f"/{TEAM_LEARNINGS}"},
        headers=auth,
    )
    assert r.status_code == 404, r.text


@pytest.mark.parametrize("method", ["put", "post"])
@pytest.mark.parametrize("path", TEAM_MEMORY_SPELLINGS)
async def test_team_memory_write_denied_409(env: Any, method: str, path: str) -> None:
    c, _srv, auth, aid = env
    r = await _team_memory_write(c, auth, aid, method=method, path=path, locale="en")
    assert r.status_code == 409, r.text
    body = r.json()["error"]
    assert body["code"] == "TEAM_ARTIFACT_OWNERSHIP_DENIED"
    assert body["message"] == TEAM_MEMORY_DENIED_EN
    await _assert_not_on_disk(c, auth, aid)


@pytest.mark.parametrize("method", ["put", "post"])
async def test_team_memory_write_denied_copy_is_localized(env: Any, method: str) -> None:
    c, _srv, auth, aid = env
    r = await _team_memory_write(
        c, auth, aid, method=method, path=f"/{TEAM_LEARNINGS}", locale="zh"
    )
    assert r.status_code == 409, r.text
    assert r.json()["error"]["message"] == TEAM_MEMORY_DENIED_ZH


async def _seed_team_memory(c: Any, auth: dict[str, str], aid: str, blob: bytes) -> None:
    """Seed via the sanctioned writer (archive restore) so the bytes are observable."""
    r = await c.post(
        f"/api/agents/{aid}/workspace/archive",
        params={"mode": "merge"},
        headers=auth,
        files={"file": ("workspace.zip", _zip_with({TEAM_LEARNINGS: blob}), "application/zip")},
    )
    assert r.status_code == 200, r.text


async def _read_team_memory(c: Any, auth: dict[str, str], aid: str) -> Any:
    return await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": f"/{TEAM_LEARNINGS}"},
        headers=auth,
    )


def _default_param_team_spellings(srv: Any, aid: str) -> dict[str, str]:
    """Three spellings that reach ``.octop/team/LEARNINGS.md`` for the same agent.

    ``host_absolute`` is the FIND-20 hole: with ``from_workspace`` omitted (default
    ``false``) a leading ``/`` is a **host** path, so the guard used to see
    ``Users/…/.octop/team/…`` and miss the protected surface. ``file_url`` is the
    FIND-21 spelling, which always resolves host-absolute.
    """
    host = resolve_agent_workspace_dir(srv, aid) / TEAM_LEARNINGS
    return {
        "relative": TEAM_LEARNINGS,
        "host_absolute": host.as_posix(),
        "file_url": host.as_uri(),
    }


async def _default_param_team_write(
    c: Any, auth: dict[str, str], aid: str, *, method: str, path: str
) -> Any:
    """``from_workspace`` is deliberately **omitted** — that is the defect under test."""
    headers = {**auth, "Accept-Language": "en"}
    if method == "put":
        return await c.put(
            f"/api/agents/{aid}/workspace/file",
            params={"path": path},
            headers=headers,
            json={"content": "forged\n"},
        )
    return await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={"path": path},
        headers=headers,
        files={"file": ("LEARNINGS.md", b"forged\n", "text/markdown")},
    )


@pytest.mark.parametrize("method", ["put", "post"])
@pytest.mark.parametrize("shape", ["relative", "host_absolute", "file_url"])
async def test_team_memory_write_denied_with_default_params_409(
    env: Any, method: str, shape: str
) -> None:
    """Default ``from_workspace`` + 3 spellings ⇒ 409 (code + copy verbatim), nothing written."""
    c, srv, auth, aid = env
    sentinel = b"# owned by the team runtime\n"
    await _seed_team_memory(c, auth, aid, sentinel)
    path = _default_param_team_spellings(srv, aid)[shape]
    r = await _default_param_team_write(c, auth, aid, method=method, path=path)
    assert r.status_code == 409, r.text
    body = r.json()["error"]
    assert body["code"] == "TEAM_ARTIFACT_OWNERSHIP_DENIED"
    assert body["message"] == TEAM_MEMORY_DENIED_EN
    after = await _read_team_memory(c, auth, aid)
    assert after.status_code == 200, after.text
    assert after.json()["content"].encode() == sentinel  # bytes unchanged


@pytest.mark.parametrize("from_workspace_default", [True, False])
@pytest.mark.parametrize(
    "rel",
    [
        "SOUL.md",
        "outbound/a.txt",
        "skills/demo/SKILL.md",
        ".octop/sessions/state.db",  # in ``.octop`` but not ``.octop/team``
        ".octop/auth/token.json",  # in ``.octop`` but not ``.octop/team``
    ],
)
async def test_ordinary_files_still_writable_in_both_modes(
    env: Any, rel: str, from_workspace_default: bool
) -> None:
    """Positive control: the widened guard must not touch ordinary or ``.octop`` siblings."""
    c, _srv, auth, aid = env
    path = f"/{rel}" if from_workspace_default else rel
    params = {"path": path}
    if from_workspace_default:
        params["from_workspace"] = "true"
    r = await c.put(
        f"/api/agents/{aid}/workspace/file",
        params=params,
        headers=auth,
        json={"content": "ok\n"},
    )
    assert r.status_code == 200, r.text
    back = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": f"/{rel}"},
        headers=auth,
    )
    assert back.status_code == 200, back.text
    assert back.json()["content"] == "ok\n"


async def test_archive_import_still_writes_team_memory(env: Any) -> None:
    """Positive control: the registered archive-restore exemption keeps writing ``.octop/**``."""
    c, _srv, auth, aid = env
    r = await c.post(
        f"/api/agents/{aid}/workspace/archive",
        params={"mode": "merge"},
        headers=auth,
        files={
            "file": (
                "workspace.zip",
                _zip_with({TEAM_LEARNINGS: b"# restored\n"}),
                "application/zip",
            )
        },
    )
    assert r.status_code == 200, r.text
    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": f"/{TEAM_LEARNINGS}"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "# restored\n"


@pytest.mark.parametrize(
    "path",
    [
        "/./_builtin_skills/foo/SKILL.md",
        "/x/../_builtin_skills/foo/SKILL.md",
        "//_builtin_skills/foo/SKILL.md",
    ],
)
async def test_builtin_skills_bypass_spellings_forbidden(env: Any, path: str) -> None:
    """SEC-10: the shared normalizer closes these bypasses on the pre-existing 403 guard.

    ``delete`` is the endpoint that already carried ``_assert_workspace_mutable``
    (``write_file`` / ``upload_file`` never guarded ``_builtin_skills`` at all).
    """
    c, _srv, auth, aid = env
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": path},
        headers=auth,
    )
    assert r.status_code == 403, r.text
